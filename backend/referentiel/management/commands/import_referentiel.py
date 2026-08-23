import os
import openpyxl
from django.core.management.base import BaseCommand
from django.conf import settings
from django.db import transaction

from referentiel.models import Referentiel, Theme, Question, ChoixReponse


class Command(BaseCommand):
    help = "Importe le référentiel (thèmes, questions, choix) depuis le fichier Excel."

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            type=str,
            default=os.path.join(
                settings.BASE_DIR, "referentiel", "data", "referentiel_29_questions.xlsx"
            ),
            help="Chemin du fichier Excel à importer.",
        )

    def handle(self, *args, **options):
        filepath = options["file"]

        if not os.path.exists(filepath):
            self.stderr.write(self.style.ERROR(f"Fichier introuvable : {filepath}"))
            return

        wb = openpyxl.load_workbook(filepath, data_only=True)

        with transaction.atomic():
            referentiel = self._import_referentiel(wb["Referentiel"])
            themes_map = self._import_themes(wb["Themes"], referentiel)
            questions_map = self._import_questions(wb["Questions"], themes_map)
            self._import_choix(wb["Choix"], questions_map)

        self.stdout.write(self.style.SUCCESS("Import terminé avec succès."))

    # -----------------------------------------------------------------
    def _import_referentiel(self, ws):
        row = next(ws.iter_rows(min_row=2, values_only=True))
        nom, version, description, score_maximum = row

        referentiel, created = Referentiel.objects.update_or_create(
            nom=nom,
            version=version,
            defaults={
                "description": description or "",
                "score_maximum": score_maximum,
                "actif": True,
            },
        )
        action = "créé" if created else "mis à jour"
        self.stdout.write(f"Référentiel {action} : {referentiel}")
        return referentiel

    # -----------------------------------------------------------------
    def _import_themes(self, ws, referentiel):
        themes_map = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            code, nom, description, ordre = row
            if code is None:
                continue
            theme, created = Theme.objects.update_or_create(
                referentiel=referentiel,
                code=code,
                defaults={
                    "nom": nom,
                    "description": description or "",
                    "ordre": ordre,
                },
            )
            themes_map[code] = theme
            action = "créé" if created else "mis à jour"
            self.stdout.write(f"  Thème {action} : {theme}")
        return themes_map

    # -----------------------------------------------------------------
    def _import_questions(self, ws, themes_map):
        questions_map = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            theme_code, numero, texte, ref_iso, ref_nist, ref_loi, recommandation, ordre = row
            if numero is None:
                continue

            theme = themes_map.get(theme_code)
            if theme is None:
                self.stderr.write(
                    self.style.WARNING(f"  Thème inconnu '{theme_code}' pour la question {numero}, ignorée.")
                )
                continue

            question, created = Question.objects.update_or_create(
                theme=theme,
                numero=numero,
                defaults={
                    "texte": texte,
                    "ref_iso27001": ref_iso or "",
                    "ref_nist": ref_nist or "",
                    "ref_loi0908": ref_loi or "",
                    "recommandation": recommandation or "",
                    "ordre": ordre,
                },
            )
            questions_map[numero] = question
            action = "créée" if created else "mise à jour"
            self.stdout.write(f"    Question {action} : Q{numero}")
        return questions_map

    # -----------------------------------------------------------------
    def _import_choix(self, ws, questions_map):
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            question_numero, valeur, libelle_niveau, texte = row
            if question_numero is None:
                continue

            question = questions_map.get(question_numero)
            if question is None:
                self.stderr.write(
                    self.style.WARNING(f"  Question inconnue Q{question_numero}, choix ignoré.")
                )
                continue

            ChoixReponse.objects.update_or_create(
                question=question,
                valeur=valeur,
                defaults={
                    "libelle_niveau": libelle_niveau,
                    "texte": texte,
                },
            )
            count += 1
        self.stdout.write(f"      {count} choix importés/mis à jour.")